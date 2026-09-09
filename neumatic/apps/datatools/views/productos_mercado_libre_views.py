# neumatic\apps\datatools\views\productos_mercado_libre_views.py
from django.views.generic import TemplateView
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.db.models import Q, Sum
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib import messages
from django.shortcuts import redirect
from django.urls import reverse
from django.http import JsonResponse, HttpResponse
from datetime import datetime
import logging
import csv
import json
import decimal
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from io import BytesIO, StringIO

from apps.maestros.models.base_models import ProductoEstado
from apps.maestros.models.producto_models import Producto
from apps.datatools.forms.productos_mercado_libre_forms import (
	ExportarProductosForm,
	ExportarProductosArchivoForm
)


logger = logging.getLogger(__name__)


class ExportarProductosCarritoView(LoginRequiredMixin, TemplateView):
	template_name = 'datatools/productos_mercado_libre_form.html'
	paginate_by = 20
	
	def get_queryset(self, form):
		"""Obtiene el queryset aplicando todos los filtros"""
		productos = Producto.objects.filter(estatus_producto=True).select_related(
			'id_marca', 'id_familia', 'id_modelo', 'id_cai'
		)
		
		#-- Aplicar filtros.
		tipo_producto = form.cleaned_data.get('tipo_producto')
		if tipo_producto:
			productos = productos.filter(tipo_producto=tipo_producto)
		
		marca = form.cleaned_data.get('marca')
		if marca:
			productos = productos.filter(id_marca=marca)
		
		familia = form.cleaned_data.get('familia')
		if familia:
			productos = productos.filter(id_familia=familia)
		
		modelo = form.cleaned_data.get('modelo')
		if modelo:
			productos = productos.filter(id_modelo=modelo)
		
		medida = form.cleaned_data.get('medida')
		if medida:
			productos = productos.filter(medida__icontains=medida)
		
		cai = form.cleaned_data.get('cai')
		if cai:
			productos = productos.filter(id_cai__cai__icontains=cai)
		
		busqueda = form.cleaned_data.get('busqueda')
		if busqueda:
			productos = productos.filter(
				Q(nombre_producto__icontains=busqueda) |
				Q(descripcion_producto__icontains=busqueda)
			)
		
		#-- Anotar el stock total para todos los productos.
		productos = productos.annotate(
			stock_total=Sum('productostock__stock')
		)
		
		#-- Aplicar filtro solo_con_stock (si está activo).
		solo_con_stock = form.cleaned_data.get('solo_con_stock')
		if solo_con_stock:
			productos = productos.filter(stock_total__gt=0)
		
		carrito = form.cleaned_data.get('carrito')
		if carrito == 'si':
			productos = productos.filter(carrito=True)
		elif carrito == 'no':
			productos = productos.filter(carrito=False)
		
		return productos.order_by('nombre_producto')
	
	def get(self, request, *args, **kwargs):
		context = self.get_context_data()
		form = ExportarProductosForm(request.GET)
		context['form'] = form
		context['titulo'] = 'Gestión de productos (E-Commerce)'
		
		tiene_filtros = any(key for key in request.GET.keys() if key not in ['page', 'accion'])
		productos = Producto.objects.none()
		total_registros = 0
		todos_ids = []
		
		if tiene_filtros and form.is_valid():
			productos = self.get_queryset(form)
			total_registros = productos.count()
			
			#-- Obtener TODOS los IDs filtrados (sin paginación).
			todos_ids = list(productos.values_list('id_producto', flat=True))
			
			#-- Paginación.
			paginator = Paginator(productos, self.paginate_by)
			page = request.GET.get('page', 1)
			
			try:
				page_obj = paginator.page(page)
			except PageNotAnInteger:
				page_obj = paginator.page(1)
			except EmptyPage:
				page_obj = paginator.page(paginator.num_pages)
			
			context.update({
				'page_obj': page_obj,
				'productos': page_obj.object_list,
				'total_registros': total_registros,
				'total_paginas': paginator.num_pages,
				'rango_paginas': self.get_page_range(page_obj, paginator),
				'todos_ids': todos_ids,
			})
		else:
			context.update({
				'page_obj': None,
				'productos': [],
				'total_registros': 0,
				'total_paginas': 0,
				'rango_paginas': [],
				'todos_ids': [],
			})
		
		#-- Procesar acción masiva si se envió.
		accion = request.GET.get('accion')
		if accion and hasattr(self, f'procesar_{accion}'):
			#-- Obtener los IDs de productos seleccionados de la URL.
			productos_ids_str = request.GET.get('productos_ids', '')
			productos_ids = []
			if productos_ids_str:
				try:
					productos_ids = [int(id) for id in productos_ids_str.split(',') if id.strip()]
				except ValueError:
					pass
			
			#-- Si no hay IDs seleccionados, mostrar mensaje de advertencia.
			if not productos_ids:
				messages.warning(request, 'No hay productos seleccionados para procesar.')
				url = reverse('exportar_productos_carrito')
				query_params = request.GET.copy()
				query_params.pop('accion', None)
				query_params.pop('productos_ids', None)
				return redirect(f'{url}?{query_params.urlencode()}')
			
			#-- Filtrar el queryset para solo los productos seleccionados.
			productos_seleccionados = Producto.objects.filter(
				id_producto__in=productos_ids,
				estatus_producto=True
			)
			
			#-- Si no hay productos válidos, mostrar mensaje.
			if not productos_seleccionados.exists():
				messages.warning(request, 'No se encontraron productos válidos para procesar.')
				url = reverse('exportar_productos_carrito')
				query_params = request.GET.copy()
				query_params.pop('accion', None)
				query_params.pop('productos_ids', None)
				return redirect(f'{url}?{query_params.urlencode()}')
			
			return getattr(self, f'procesar_{accion}')(request, form, productos_seleccionados)
		
		return self.render_to_response(context)
	
	def get_page_range(self, page_obj, paginator):
		"""Obtiene el rango de páginas para mostrar (máximo 5)"""
		current = page_obj.number
		total = paginator.num_pages
		
		if total <= 5:
			return range(1, total + 1)
		
		if current <= 3:
			return range(1, 6)
		elif current >= total - 2:
			return range(total - 4, total + 1)
		else:
			return range(current - 2, current + 3)
	
	def procesar_activar_carrito(self, request, form, productos):
		"""Activa el campo carrito para los productos seleccionados"""
		return self._procesar_accion_masiva(request, productos, True)
	
	def procesar_desactivar_carrito(self, request, form, productos):
		"""Desactiva el campo carrito para los productos seleccionados"""
		return self._procesar_accion_masiva(request, productos, False)
	
	def procesar_toggle_carrito(self, request, form, productos):
		"""Invierte el campo carrito para los productos seleccionados"""
		productos_ids = list(productos.values_list('id_producto', flat=True))
		
		if not productos_ids:
			messages.warning(request, 'No hay productos seleccionados para procesar.')
			url = reverse('exportar_productos_carrito')
			query_params = request.GET.copy()
			query_params.pop('accion', None)
			query_params.pop('productos_ids', None)
			return redirect(f'{url}?{query_params.urlencode()}')
		
		#-- Invertir estado SOLO para los productos seleccionados.
		for producto in productos:
			producto.carrito = not producto.carrito
		
		Producto.objects.bulk_update(productos, ['carrito'])
		
		cantidad_actualizados = productos.count()
		messages.success(
			request,
			f'Se invirtió el estado "Carrito" para {cantidad_actualizados} productos.'
		)
		
		url = reverse('exportar_productos_carrito')
		query_params = request.GET.copy()
		query_params.pop('accion', None)
		query_params.pop('productos_ids', None)
		return redirect(f'{url}?{query_params.urlencode()}')
	
	def _procesar_accion_masiva(self, request, productos, valor):
		"""Procesa la acción masiva para activar/desactivar carrito SOLO para seleccionados"""
		productos_ids = list(productos.values_list('id_producto', flat=True))
		
		if not productos_ids:
			messages.warning(request, 'No hay productos seleccionados para procesar.')
			url = reverse('exportar_productos_carrito')
			query_params = request.GET.copy()
			query_params.pop('accion', None)
			query_params.pop('productos_ids', None)
			return redirect(f'{url}?{query_params.urlencode()}')
		
		cantidad_actualizados = Producto.objects.filter(
			id_producto__in=productos_ids
		).update(carrito=valor)
		
		estado_texto = 'activado' if valor else 'desactivado'
		messages.success(
			request,
			f'Se {estado_texto} el campo "Carrito" para {cantidad_actualizados} productos.'
		)
		
		url = reverse('exportar_productos_carrito')
		query_params = request.GET.copy()
		query_params.pop('accion', None)
		query_params.pop('productos_ids', None)
		return redirect(f'{url}?{query_params.urlencode()}')
	
	def post(self, request, *args, **kwargs):
		"""Maneja peticiones POST para actualizar carrito individual"""
		try:
			producto_id = request.POST.get('producto_id')
			valor = request.POST.get('valor')
			
			if not producto_id or valor is None:
				return JsonResponse({
					'success': False,
					'error': 'Datos incompletos'
				}, status=400)
			
			producto = Producto.objects.get(id_producto=producto_id, estatus_producto=True)
			nuevo_valor = valor.lower() == 'true'
			producto.carrito = nuevo_valor
			producto.save(update_fields=['carrito'])
			
			return JsonResponse({
				'success': True,
				'message': 'Estado actualizado correctamente'
			})
			
		except Producto.DoesNotExist:
			return JsonResponse({
				'success': False,
				'error': 'Producto no encontrado'
			}, status=404)
		except Exception as e:
			logger.error(f'Error actualizando carrito: {str(e)}')
			return JsonResponse({
				'success': False,
				'error': 'Error al actualizar el estado'
			}, status=500)


class ExportarProductosArchivoView(LoginRequiredMixin, TemplateView):
	"""Vista para exportar productos a archivo"""
	
	template_name = 'datatools/exportar_productos_archivo.html'
	
	def get_queryset_export(self, estados_seleccionados):
		"""
		Obtiene el queryset con los datos para exportar
		"""
		#-- Filtrar productos activos y con carrito=True.
		productos = Producto.objects.filter(
			estatus_producto=True,
			carrito=True
		)
		
		#-- Filtrar por estados seleccionados.
		if estados_seleccionados:
			codigos_estado = ProductoEstado.objects.filter(
				id_producto_estado__in=estados_seleccionados
			).values_list('estado_producto', flat=True)
			
			productos = productos.filter(
				id_producto_estado__estado_producto__in=codigos_estado
			)
		
		#-- Anotar el stock total.
		productos = productos.annotate(
			stock_total=Sum('productostock__stock')
		)
		
		#-- Seleccionar los campos necesarios.
		productos = productos.values(
			'id_cai__cai',
			'nombre_producto',
			'medida',
			'stock_total',
			'descuento',
			'precio',
			'id_producto_estado__estado_producto'
		).order_by('id_cai__cai')
		
		return productos
	
	def get(self, request, *args, **kwargs):
		context = self.get_context_data()
		form = ExportarProductosArchivoForm()
		context['form'] = form
		context['titulo'] = 'Exportar productos (E-Commerce)'
		
		#-- Si hay parámetros GET, procesar.
		if request.GET.get('exportar'):
			form = ExportarProductosArchivoForm(request.GET)
			if form.is_valid():
				return self.procesar_exportacion(request, form)
			else:
				#-- Si hay errores, mostrar mensajes.
				for field, errors in form.errors.items():
					for error in errors:
						messages.error(request, f'{field}: {error}')
				context['form'] = form
		
		return self.render_to_response(context)
	
	def post(self, request, *args, **kwargs):
		"""Maneja la exportación desde un formulario POST"""
		form = ExportarProductosArchivoForm(request.POST)
		if form.is_valid():
			return self.procesar_exportacion(request, form)
		else:
			#-- Si hay errores, mostrar mensajes y volver al formulario.
			for field, errors in form.errors.items():
				for error in errors:
					messages.error(request, f'{field}: {error}')
			
			context = self.get_context_data()
			context['form'] = form
			return self.render_to_response(context)
	
	def procesar_exportacion(self, request, form):
		"""Procesa la exportación de datos"""
		
		#-- Obtener datos del formulario.
		estados_ids = form.cleaned_data.get('estados')
		formato = form.cleaned_data.get('formato')
		separador_decimal = form.cleaned_data.get('separador_decimal')
		
		#-- Obtener los datos.
		datos = self.get_queryset_export(estados_ids)
		datos_lista = list(datos)
		
		if not datos_lista:
			messages.warning(request, 'No hay datos para exportar con los filtros seleccionados.')
			return redirect('exportar_productos_archivo')
		
		#-- Generar nombre de archivo con timestamp.
		timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
		nombre_base = f"productos_carrito_{timestamp}"
		
		#-- Preparar datos para exportación (CONVERTIR DECIMAL A FLOAT).
		datos_export = []
		for item in datos_lista:
			row = dict(item)
			
			#-- Obtener valores y convertir Decimal a float.
			precio = row.get('precio')
			descuento = row.get('descuento')
			stock = row.get('stock_total')
			
			#-- Convertir Decimal a float si es necesario
			if isinstance(precio, decimal.Decimal):
				precio = float(precio)
			if isinstance(descuento, decimal.Decimal):
				descuento = float(descuento)
			if isinstance(stock, decimal.Decimal):
				stock = float(stock)
			
			#-- Asegurar que no sean None
			precio = precio or 0.0
			descuento = descuento or 0.0
			stock = stock or 0.0
			
			#-- Calcular oferta
			oferta = round(precio - (precio * descuento / 100), 2)
			
			#-- Obtener CAI y Estado (pueden ser None)
			cai = row.get('id_cai__cai') or ''
			estado = row.get('id_producto_estado__estado_producto') or ''
			
			row_export = {
				'CAI': cai,
				'Producto': row.get('nombre_producto') or '',
				'Medida': row.get('medida') or '',
				'Stock': stock,
				'Descuento (%)': descuento,
				'Oferta': oferta,
				'Precio': precio,
				'Estado': estado,
			}
			datos_export.append(row_export)
		
		#-- Exportar según formato.
		if formato == 'csv':
			return self.exportar_csv_descarga(request, datos_export, nombre_base, separador_decimal)
		elif formato == 'xlsx':
			return self.exportar_excel_descarga(request, datos_export, nombre_base)
		elif formato == 'json':
			return self.exportar_json_descarga(request, datos_export, nombre_base)
		elif formato == 'txt':
			return self.exportar_txt_descarga(request, datos_export, nombre_base)
		
		messages.error(request, 'Formato de exportación no válido.')
		return redirect('exportar_productos_archivo')
	
	def exportar_csv_descarga(self, request, datos, nombre_base, separador_decimal):
		"""Exporta los datos a CSV y lo envía como descarga directa"""
		
		if not datos:
			messages.warning(request, 'No hay datos para exportar.')
			return redirect('exportar_productos_archivo')
		
		#-- Crear archivo CSV en memoria.
		output = StringIO()
		headers = list(datos[0].keys())
		
		writer = csv.DictWriter(output, fieldnames=headers, delimiter=';')
		writer.writeheader()
		
		#-- Escribir datos con formato numérico.
		for row in datos:
			row_formateada = {}
			for key, value in row.items():
				if isinstance(value, (int, float)):
					#-- Formatear numéricos con 2 decimales
					if separador_decimal == 'coma':
						row_formateada[key] = f"{value:.2f}".replace('.', ',')
					else:
						row_formateada[key] = f"{value:.2f}"
				else:
					row_formateada[key] = value
			writer.writerow(row_formateada)
		
		#-- Preparar respuesta HTTP.
		nombre_archivo = f"{nombre_base}.csv"
		response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8-sig')
		response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
		
		return response
	
	def exportar_excel_descarga(self, request, datos, nombre_base):
		"""Exporta los datos a Excel y lo envía como descarga directa"""
		
		if not datos:
			messages.warning(request, 'No hay datos para exportar.')
			return redirect('exportar_productos_archivo')
		
		#-- Crear libro de Excel en memoria.
		wb = openpyxl.Workbook()
		ws = wb.active
		ws.title = "Productos Carrito"
		
		#-- Estilos.
		header_font = Font(bold=True, color="FFFFFF")
		header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
		header_alignment = Alignment(horizontal="center", vertical="center")
		numeric_alignment = Alignment(horizontal="right", vertical="center")
		text_alignment = Alignment(horizontal="left", vertical="center")
		
		#-- Escribir encabezados.
		headers = list(datos[0].keys())
		for col_idx, header in enumerate(headers, 1):
			cell = ws.cell(row=1, column=col_idx, value=header)
			cell.font = header_font
			cell.fill = header_fill
			cell.alignment = header_alignment
		
		#-- Identificar columnas numéricas y sus tipos.
		columnas_decimales = ['Descuento (%)', 'Oferta', 'Precio']
		columna_entera = 'Stock'
		
		#-- Escribir datos.
		for row_idx, row_data in enumerate(datos, 2):
			for col_idx, header in enumerate(headers, 1):
				value = row_data.get(header, '')
				cell = ws.cell(row=row_idx, column=col_idx, value=value)
				
				#-- Aplicar alineación y formato según tipo.
				if isinstance(value, (int, float)):
					cell.alignment = numeric_alignment
					
					#-- Aplicar formato específico según la columna.
					if header == columna_entera:
						#-- Formato de número entero (sin decimales).
						cell.number_format = '#,##0'
					elif header in columnas_decimales:
						#-- Formato de número con 2 decimales.
						cell.number_format = '#,##0.00'
					else:
						#-- Formato genérico para otros numéricos.
						cell.number_format = '#,##0.00'
				else:
					cell.alignment = text_alignment
		
		#-- Ajustar ancho de columnas.
		for col_idx, header in enumerate(headers, 1):
			max_length = len(header)
			for row in datos:
				value = str(row.get(header, ''))
				if len(value) > max_length:
					max_length = len(value)
			adjusted_width = min(max_length + 2, 50)
			ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
		
		#-- Guardar en BytesIO.
		output = BytesIO()
		wb.save(output)
		output.seek(0)
		
		#-- Preparar respuesta HTTP.
		nombre_archivo = f"{nombre_base}.xlsx"
		response = HttpResponse(
			output.getvalue(),
			content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
		)
		response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
		
		return response
	
	def exportar_json_descarga(self, request, datos, nombre_base):
		"""Exporta los datos a JSON y lo envía como descarga directa"""
		
		if not datos:
			messages.warning(request, 'No hay datos para exportar.')
			return redirect('exportar_productos_archivo')
		
		#-- Los datos ya vienen con Decimal convertido a float desde procesar_exportacion
		#-- Solo redondear floats a 2 decimales.
		datos_formateados = []
		for row in datos:
			row_formateada = {}
			for key, value in row.items():
				if isinstance(value, float):
					row_formateada[key] = round(value, 2)
				else:
					row_formateada[key] = value
			datos_formateados.append(row_formateada)
		
		#-- Generar JSON.
		json_content = json.dumps(datos_formateados, ensure_ascii=False, indent=2)
		
		#-- Preparar respuesta HTTP.
		nombre_archivo = f"{nombre_base}.json"
		response = HttpResponse(json_content, content_type='application/json; charset=utf-8')
		response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
		
		return response
	
	def exportar_txt_descarga(self, request, datos, nombre_base):
		"""Exporta los datos a TXT y lo envía como descarga directa"""
		
		if not datos:
			messages.warning(request, 'No hay datos para exportar.')
			return redirect('exportar_productos_archivo')
		
		headers = list(datos[0].keys())
		
		#-- Detectar columnas numéricas.
		columnas_numericas = []
		for header in headers:
			es_numerica = True
			max_muestras = min(10, len(datos))
			
			for i, row in enumerate(datos):
				if i >= max_muestras:
					break
				valor = row.get(header, '')
				if valor is None or valor == '':
					continue
				if isinstance(valor, (int, float)):
					continue
				valor_str = str(valor)
				valor_limpio = valor_str.replace('.', '').replace(',', '').strip()
				try:
					float(valor_limpio)
				except (ValueError, TypeError):
					es_numerica = False
					break
			
			if es_numerica:
				columnas_numericas.append(header)
		
		#-- Calcular anchos de columna.
		col_widths = {}
		for header in headers:
			max_len = len(header)
			for row in datos:
				value = row.get(header, '')
				
				#-- Si es numérico, formatear a string con 2 decimales para calcular ancho.
				if header in columnas_numericas:
					try:
						value = f"{float(value):.2f}"
					except (ValueError, TypeError):
						value = str(value)
				else:
					value = str(value)
				
				if len(value) > max_len:
					max_len = len(value)
			col_widths[header] = max_len + 2
		
		#-- Generar contenido TXT.
		output = StringIO()
		total_width = sum(col_widths.values()) + len(headers) + 12
		output.write('=' * total_width + '\n')
		
		#-- Encabezados.
		header_line = ' | '.join([header.ljust(col_widths[header]) for header in headers])
		output.write(header_line + '\n')
		output.write('-' * total_width + '\n')
		
		#-- Datos.
		for row in datos:
			line_parts = []
			for header in headers:
				value = str(row.get(header, ''))
				
				#-- Si es numérico, formatear con 2 decimales.
				if header in columnas_numericas:
					try:
						#-- Intentar convertir a float y formatear con 2 decimales.
						value = f"{float(value):.2f}"
					except (ValueError, TypeError):
						#-- Si falla, dejar como string.
						value = str(value)
				else:
					value = str(value)
				
				#-- Alinear a la derecha si es numérico, izquierda si es texto.
				if header in columnas_numericas:
					line_parts.append(value.rjust(col_widths[header]))
				else:
					line_parts.append(value.ljust(col_widths[header]))
			line = ' | '.join(line_parts)
			output.write(line + '\n')
		
		output.write('=' * total_width + '\n')
		output.write(f'\nTotal de registros: {len(datos)}\n')
		output.write(f'Fecha de exportación: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}\n')
		
		#-- Preparar respuesta HTTP.
		nombre_archivo = f"{nombre_base}.txt"
		response = HttpResponse(output.getvalue(), content_type='text/plain; charset=utf-8')
		response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
		
		return response